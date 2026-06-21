"""Export draw.io diagrams and insert PNG images into SmartMed Report.docx."""
import os
import shutil
import subprocess
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

BASE = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
REPORT_PATH = os.path.join(BASE, "Report", "SmartMed Report.docx")
DIAGRAMS_DIR = os.path.join(BASE, "Diagrams")
EXPORT_DIR = os.path.join(DIAGRAMS_DIR, "export")
MOCKS_DIR = os.path.join(BASE, "DesignMocks")

DIAGRAM_HEADINGS = [
    ("Architecture Diagram", "Architecture.png"),
    ("Use Case Diagram", "UseCase.png"),
    ("ER Diagram", "ER.png"),
    ("Class Diagram", "Class.png"),
    ("Sequence Diagram", "Sequence.png"),
    ("Database Design", "ER.png"),
]

OUTPUT_SCREENS = [
    ("Login screen", "login_screen/screen.png"),
    ("Admin dashboard", "admin_dashboard/screen.png"),
    ("Manage medicines", "manage_medicines/screen.png"),
    ("Manage customers", "manage_customers/screen.png"),
    ("Manage orders", "manage_orders/screen.png"),
    ("Generate reports", "generate_reports/screen.png"),
]

DRAWIO_CANDIDATES = [
    os.path.join(os.environ.get("TEMP", ""), "drawio-export", "draw.io.exe"),
    os.path.join(os.environ.get("ProgramFiles", ""), "draw.io", "draw.io.exe"),
    os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "draw.io", "draw.io.exe"),
]


def find_drawio():
    for path in DRAWIO_CANDIDATES:
        if path and os.path.isfile(path):
            return path
    return None


def export_diagrams(drawio_exe):
    os.makedirs(EXPORT_DIR, exist_ok=True)
    files = {
        "Architecture.drawio": "Architecture.png",
        "Usecase.drawio": "UseCase.png",
        "ER.drawio": "ER.png",
        "Class.drawio": "Class.png",
        "Sequence.drawio": "Sequence.png",
    }
    for src, dst in files.items():
        out = os.path.join(EXPORT_DIR, dst)
        if os.path.isfile(out):
            print(f"Using existing {dst}")
            continue
        inp = os.path.join(DIAGRAMS_DIR, src)
        if not os.path.isfile(inp):
            print(f"Missing source: {inp}")
            continue
        cmd = [drawio_exe, "-x", "-f", "png", "-b", "10", "--width", "900", "-o", out, inp]
        try:
            subprocess.run(cmd, check=False, capture_output=True, timeout=60)
        except subprocess.TimeoutExpired:
            print(f"Export timed out: {dst}")
            continue
        if os.path.isfile(out):
            print(f"Exported {dst}")
        else:
            print(f"Export failed: {dst}")


def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def paragraph_has_image(paragraph):
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            return True
    return False


def section_has_image(doc, heading_text):
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() != heading_text or not para.style.name.startswith("Heading"):
            continue
        for j in range(i + 1, min(i + 5, len(doc.paragraphs))):
            if doc.paragraphs[j].style.name.startswith("Heading"):
                break
            if paragraph_has_image(doc.paragraphs[j]):
                return True
    return False


def insert_image_after_body(doc, heading_text, image_path, caption=None, width=6.0):
    if not os.path.isfile(image_path):
        print(f"Image not found: {image_path}")
        return False

    if section_has_image(doc, heading_text):
        print(f"Already inserted: {heading_text}")
        return True

    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() != heading_text or not para.style.name.startswith("Heading"):
            continue

        anchor = doc.paragraphs[i + 1] if i + 1 < len(doc.paragraphs) else para
        img_p = insert_paragraph_after(anchor)
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.add_run().add_picture(image_path, width=Inches(width))

        if caption:
            cap = insert_paragraph_after(img_p)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = cap.add_run(caption)
            cap_run.font.name = "Times New Roman"
            cap_run.font.size = Pt(10)
            cap_run.italic = True

        print(f"Inserted image for: {heading_text}")
        return True

    print(f"Heading not found: {heading_text}")
    return False


def insert_output_screens(doc):
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() != "Output Screens" or not para.style.name.startswith("Heading"):
            continue

        anchor = para
        for label, rel in OUTPUT_SCREENS:
            img = os.path.join(MOCKS_DIR, rel.replace("/", os.sep))
            if not os.path.isfile(img):
                print(f"Mock not found: {img}")
                continue

            cap_p = insert_paragraph_after(anchor)
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = cap_p.add_run(label)
            cap_run.font.name = "Times New Roman"
            cap_run.font.size = Pt(11)
            cap_run.bold = True

            img_p = insert_paragraph_after(cap_p)
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.add_run().add_picture(img, width=Inches(5.5))

            anchor = img_p
            print(f"Inserted output screen: {label}")

        # Remove placeholder lines after screens were added
        j = i + 1
        while j < len(doc.paragraphs):
            text = doc.paragraphs[j].text.strip()
            if doc.paragraphs[j].style.name.startswith("Heading"):
                break
            if text.startswith("[Insert screenshot:") or text.startswith("Design mock PNG"):
                rm = doc.paragraphs[j]
                rm._element.getparent().remove(rm._element)
                continue
            j += 1
        return True
    return False


def main():
    drawio = find_drawio()
    if drawio:
        export_diagrams(drawio)
    else:
        print("draw.io not found; using existing PNG files in Docs/Diagrams/export/")

    doc = Document(REPORT_PATH)

    for heading, filename in DIAGRAM_HEADINGS:
        path = os.path.join(EXPORT_DIR, filename)
        caption = f"Figure: {heading}"
        insert_image_after_body(doc, heading, path, caption=caption)

    insert_output_screens(doc)

    doc.save(REPORT_PATH)
    print(f"Saved {REPORT_PATH}")


if __name__ == "__main__":
    main()
