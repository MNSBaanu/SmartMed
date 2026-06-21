"""Replace Functional Requirements body in SmartMed Report.docx without rebuilding the whole document."""
import os
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

BASE = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
REPORT_PATH = os.path.join(BASE, "Report", "SmartMed Report.docx")

FUNCTIONAL_BODY = """Admin Features
• Login – Secure login for admins.
• Manage Medicine Details – Add, update, delete medicine details (name, category, dosage, price, stock, supplier).
• Manage Customer Details – View and update customer information.
• Manage Orders – View all orders, update order status (Pending, Ready for Pickup, Delivered).
• Generate Reports – Sales reports, stock reports, and customer order history.
• Dashboard – Overview of total sales, medicines in stock, and active orders.

Customer Features
• Register/Login – New user registration and login.
• Search Medicines – Search by name, category, or price range.
• Place Orders – Add medicines to cart and place orders.
• Track Orders – View status of orders.
• Profile Management – Update personal details and contact information.

Additional Features
• Apply discounts or promotions on medicines.
• Include medicine expiry tracking notifications.
• Include prescription upload functionality for certain medicines.
• Export order history to PDF or Excel."""


def set_body_format(paragraph):
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def find_heading_index(doc, title, level_prefix="Heading 2"):
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() == title and para.style.name.startswith(level_prefix):
            return i
    return None


def remove_paragraph_at(doc, index):
    el = doc.paragraphs[index]._element
    el.getparent().remove(el)


def insert_body_after(doc, heading_index, text):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    anchor = doc.paragraphs[heading_index]
    for part in reversed(text.split("\n\n")):
        new_p = OxmlElement("w:p")
        anchor._p.addnext(new_p)
        para = Paragraph(new_p, anchor._parent)
        para.style = doc.styles["Body Text"] if "Body Text" in [s.name for s in doc.styles] else para.style
        para.add_run(part)
        set_body_format(para)


def main():
    doc = Document(REPORT_PATH)
    start = find_heading_index(doc, "Functional Requirements")
    end = find_heading_index(doc, "Non-Functional Requirements")
    if start is None or end is None:
        raise RuntimeError("Could not locate Functional Requirements section.")

    for _ in range(end - start - 1):
        remove_paragraph_at(doc, start + 1)

    insert_body_after(doc, start, FUNCTIONAL_BODY)
    doc.save(REPORT_PATH)
    print(f"Updated Functional Requirements in {REPORT_PATH}")


if __name__ == "__main__":
    main()
