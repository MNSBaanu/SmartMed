"""Fill SmartMed Report Sample.docx — coursework report content."""
from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

BASE = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
REPORT_PATH = os.path.join(BASE, "Report", "SmartMed Report Sample.docx")
EXPORT_DIR = os.path.join(BASE, "Diagrams", "export")

# --- Functional Requirements: kept exactly as in the sample template ---
FUNCTIONAL_ADMIN = [
    "Access the system using secure administrator credentials to perform and manage all pharmacy-related operations.",
    "Maintain medicine information by adding new medicines, updating existing records, removing discontinued medicines, and managing details such as name, category, dosage, price, stock quantity, supplier information, expiry dates, and prescription requirements.",
    "Monitor and manage customer information by viewing customer records and updating customer details when necessary.",
    "Oversee customer orders by viewing all orders placed through the system and updating order statuses from Pending to Ready for Pickup or Delivered.",
    "Monitor pharmacy performance through a dashboard that provides an overview of total sales, available medicine stock, active orders, and other key operational statistics.",
    "Generate and review reports related to sales performance, medicine inventory, and customer order history to support decision-making and operational monitoring.",
    "Create, update, and remove discounts or promotional offers that can be applied to medicines.",
    "Track medicine expiry dates and receive notifications for medicines approaching expiry in order to ensure inventory safety and compliance.",
]

FUNCTIONAL_CUSTOMER = [
    "Create a personal account and securely access the system using registered credentials.",
    "Search and browse available medicines by name, category, or price range and view detailed medicine information before purchasing.",
    "Select medicines, add them to a shopping cart, review selected items, and place orders through the system.",
    "Monitor previously placed orders and track their current status throughout the order processing lifecycle.",
    "Manage personal profile information by viewing and updating personal and contact details.",
    "Upload prescription documents when purchasing medicines that require prescription verification.",
    "View available discounts and promotional offers that can be applied to eligible medicines.",
    "Export personal order history in PDF or Excel format for future reference and record keeping.",
]

FUNCTIONAL_ADDITIONAL = [
    "Low Stock Alert System - The system will automatically monitor medicine inventory levels and notify administrators when the stock quantity of a medicine falls below a predefined threshold. This helps prevent stock shortages and ensures the availability of essential medicines.",
    "Medicine Expiry Warning Dashboard - The system will identify medicines approaching their expiry dates and display warning notifications on the administrator dashboard. This assists in inventory management and reduces the risk of dispensing expired medicines.",
    "Prescription Verification System - The system will require customers to upload a valid prescription when ordering medicines marked as prescription-required. Orders containing prescription medicines cannot be processed unless the required prescription document has been uploaded and verified.",
    "Password Change Functionality - The system will allow both administrators and customers to securely change their account passwords after successful authentication, improving account security and user management.",
    "Order Cancellation - The system will allow customers to cancel orders that are still in the Pending status. Once an order has been processed or marked as Ready for Pickup, cancellation will no longer be permitted.",
]

FILL_AFTER_HEADING: dict[str, list[str]] = {
    "Architecture Diagram": [
        "The architecture diagram was drawn at the start of the project to decide how the application should be split into clear parts before writing code. "
        "A layered structure was chosen so that screens, business rules, and database access do not get mixed together, which makes the system easier to build, "
        "test, and explain in a coursework report. This approach is common in desktop business applications and matches what was taught about separation of concerns.",
        "The diagram also justified why SmartMed has separate folders for Forms, Services, Data, and Models. Each layer has one main job: the user interface "
        "handles clicks and display, services apply pharmacy rules, repositories talk to SQL Server, and model classes simply carry data between those layers. "
        "This general structure supported both the admin and customer sides of the application without duplicating database code on every screen.",
    ],
    "Use Case Diagram": [
        "The use case diagram was created during requirements analysis to confirm that both user roles — admin and customer — have a complete set of features "
        "before detailed design began. It gives a high-level, non-technical view of what the system should offer to the outside world, without showing buttons, "
        "code, or database tables.",
        "Drawing this diagram early helped justify the scope of SmartMed: admin work (stock, customers, orders, reports) and customer work (register, search, "
        "cart, track orders) are clearly separated. It also highlighted relationships such as placing an order depending on searching medicines first, which "
        "later guided how the customer screens were grouped in the application.",
    ],
    "ER Diagram": [
        "An entity-relationship diagram was prepared before writing the SQL script so that data could be stored in a logical and consistent way. "
        "The main purpose was to plan how pharmacy records link together — for example, customers place orders, and each order contains line items — "
        "so that reports and order history would still make sense after many transactions.",
        "This diagram justified the use of primary keys, foreign keys, and status rules in the database. It reduced the risk of designing tables that "
        "could not support real pharmacy operations, such as keeping past order details even when a medicine is later removed from the catalogue.",
    ],
    "Class Diagram": [
        "The class diagram was used to plan the object-oriented structure of SmartMed before implementation. It shows how domain types relate to one another "
        "through inheritance (for example, Admin and Customer extending a common User type) and how services and repositories support those entities.",
        "The diagram justified applying OOP concepts such as generalisation, encapsulation, and class responsibility at design time. It acted as a map from "
        "the problem domain to C# classes, even though the final code places most behaviour in service classes rather than inside every model class, "
        "which is explained further in the class descriptions section.",
    ],
    "Sequence Diagram": [
        "Sequence diagrams were included to show the time order of interactions for important scenarios such as login, placing an order, and managing medicines. "
        "They justify how a user action on a form flows through services and repositories before data reaches the database, which is easier to understand "
        "as a message sequence than from code alone.",
        "These diagrams were especially useful because some screens call a service layer while the overall rule remains the same: the user interface should not "
        "execute SQL directly. Comparing the sequence diagrams with the finished project helped verify that the implementation follows the intended layered flow.",
    ],
    "Database Design": [
        "The database design section documents how SmartMedDB stores pharmacy data in a reliable way. A single SQL script creates the database, tables, "
        "relationships, and sample records so the application can be demonstrated immediately after installation.",
        "A proper database layout was justified because the application must keep customer orders, stock levels, and prescription files consistent over time. "
        "Constraints on order status, foreign keys between orders and customers, and soft-delete on medicines (using an active flag) protect historical records "
        "while still allowing day-to-day catalogue changes.",
    ],
    "Installation Guide": [
        "What you need before you start",
        "• A Windows 10 or later computer",
        "• .NET Framework 4.8",
        "• Microsoft SQL Server (Express, LocalDB, or full version)",
        "• Visual Studio 2019 or later — only if building from source code",
        "",
        "Step 1 — Create the database",
        "1. Open SQL Server Management Studio.",
        "2. Connect to your SQL Server instance.",
        "3. Open Database/SmartMedDB.sql from the project folder.",
        "4. Click Execute and wait until finished.",
        "5. Check that database SmartMedDB exists with tables and sample data.",
        "",
        "Step 2 — Set the connection string",
        "1. Open SmartMed/App.config.",
        "2. Find the SmartMedDB connection string.",
        "3. Set Data Source to your SQL Server name.",
        "4. Save the file.",
        "",
        "Step 3 — Build and run",
        "1. Open SmartMed.sln in Visual Studio.",
        "2. Set SmartMed as startup project.",
        "3. Build (Ctrl+Shift+B), then run (F5).",
        "4. The login screen should open.",
        "",
        "Step 4 — Test logins",
        "• Admin: admin / admin123",
        "• Customer: customer@gmail.com / Customer123 (sample data)",
        "",
        "If you see a database warning, check Step 2 and confirm SQL Server is running.",
    ],
    "Using the Application": [],  # filled by fill_user_manual()
    "Reflection": [
        "When I started SmartMed, my main goal was to move from small classroom exercises to one complete desktop application that a pharmacy could actually use. "
        "I did not want only a login form and one table — I wanted stock, orders, customers, and reports to work together. That meant I had to plan more carefully than in earlier modules.",
        "The part I found most useful was splitting the project into folders for forms, services, data access, and models. At first I was tempted to put SQL inside model classes, like in a basic student CRUD example, "
        "but I realised that would become messy once prescriptions, promotions, and order status were added. Keeping rules in service classes made debugging easier. When checkout pricing was wrong, I could follow the flow from PlaceOrderForm to OrderService and see where the price was calculated.",
        "I also learned that design diagrams are not just for the report — they helped me before coding. The use case diagram reminded me to finish customer features and not only admin screens. The database diagram stopped me from deleting order history when a medicine is removed, which was an important business rule I might have missed.",
        "Not everything was smooth. Connecting to SQL Server on my own laptop took several attempts because the connection string in App.config had to match my instance name. WinForms layout was another challenge: admin pages are built in code for the shell form, so I had to test both the Visual Studio designer and runtime to make sure screens looked right.",
        "If I had more time, I would add automated tests for services and email alerts when order status changes. Even so, I am satisfied with the result. I understand layered design better now, and I can explain my choices clearly in a viva. Overall, SmartMed helped me connect OOP theory, database design, and real user requirements in one project.",
    ],
    "Conclusion": [
        "SmartMed is a working pharmacy management application for Windows that supports both administrator and customer roles. It covers inventory, orders, prescriptions, dashboards, and reports using C# WinForms, .NET Framework 4.8, and SQL Server.",
        "The project meets the coursework requirements through documented functional and non-functional needs, UML-style diagrams, a clear class and service structure, implementation explanation, and user instructions. Security and data integrity were considered through hashed passwords, validated inputs, and careful order and stock handling.",
        "The solution can be installed from the provided database script, demonstrated with sample logins, and extended in future work without redesigning the whole architecture. It represents a practical application of object-oriented and database concepts learned in Application Development.",
    ],
}

sys.path.insert(0, os.path.dirname(__file__))
from implementation_blocks_data import FEATURE_BLOCKS

IMPLEMENTATION_HEADING = "Brief Description of Functions"

USER_MANUAL_PARAS: list[str] = [
    "This guide walks a new user through SmartMed step by step. Follow each section in order the first time you use the system. "
    "Screenshot numbers (U1, A1, C1, etc.) match the Interface screenshot labels under each function in the Implementation section.",
    "",
    "PART 1 — Starting the application",
    "1. Double-click SmartMed.exe (or press F5 in Visual Studio after building).",
    "2. Wait for the login window to appear. If a database warning shows, contact your administrator — the app cannot save data until the database is connected.",
    "[Screenshot: U1 — Login screen]",
    "",
    "PART 2 — Common tasks (both roles)",
    "",
    "2A. Logging in as Admin",
    "1. On the login screen, select the Admin option.",
    "2. Type username: admin (or your admin username).",
    "3. Type password.",
    "4. Click Login.",
    "5. The admin home screen opens with the sidebar on the left.",
    "[Screenshot: U1 — Login as Admin]",
    "",
    "2B. Logging in as Customer",
    "1. Select Customer on the login screen.",
    "2. Enter your registered email and password.",
    "3. Click Login.",
    "4. The customer home screen opens.",
    "If you do not have an account, use Register first (section 2C).",
    "[Screenshot: U1 — Login as Customer]",
    "",
    "2C. Registering a new customer account",
    "1. From the login screen, click Register.",
    "2. Fill in full name, email, phone, address, password, and confirm password.",
    "3. Click Register / Save.",
    "4. When a success message appears, return to login and sign in with your new email and password.",
    "[Screenshot: U2 — Registration form]",
    "",
    "2D. Forgot password",
    "1. From login, click Forgot password.",
    "2. Enter the email address for your account.",
    "3. Follow the on-screen steps to set a new password.",
    "4. Return to login and sign in again.",
    "[Screenshot: U3 — Forgot password]",
    "",
    "2E. Change password (after login)",
    "1. Open Change Password from the admin or customer menu.",
    "2. Enter current password, new password, and confirm new password.",
    "3. Click Save. Use the new password next time you log in.",
    "[Screenshot: U3 — Change password]",
    "",
    "2F. Logging out",
    "1. Click Logout in the sidebar or close button on the top bar.",
    "2. You return to the login screen. Your session is cleared.",
    "",
    "PART 3 — Admin user guide",
    "",
    "3A. Admin dashboard (Overview)",
    "1. After admin login, Overview opens automatically.",
    "2. Read the summary cards: total sales, medicines in stock, active orders.",
    "3. Check expiry and low-stock alerts if shown — these warn you before products run out or expire.",
    "4. Use the sidebar to go to other admin pages.",
    "[Screenshot: A1 — Admin dashboard]",
    "",
    "3B. Manage medicines — add a new product",
    "1. Click Medicines in the sidebar.",
    "2. Click Add (or clear the form for a new entry).",
    "3. Enter medicine name, category, dosage, price, stock quantity, supplier, expiry date.",
    "4. Tick Requires Prescription if customers must upload a script.",
    "5. Set promotion dates and discount if needed.",
    "6. Click Save. The new row appears in the grid.",
    "[Screenshot: A2 — Manage medicines — add]",
    "",
    "3C. Manage medicines — edit or deactivate",
    "1. Click a row in the medicine grid.",
    "2. Change fields in the form on the right (or below).",
    "3. Click Update to save changes.",
    "4. To remove from the shop (without deleting history), click Deactivate. Confirm when asked.",
    "5. If the system blocks deactivation, an open order still references that medicine — finish or cancel that order first.",
    "[Screenshot: A2 — Manage medicines — edit/deactivate]",
    "",
    "3D. Manage customers",
    "1. Click Customers in the sidebar.",
    "2. Use search to find a customer.",
    "3. Click a row to load details.",
    "4. Edit name, phone, or address and click Save.",
    "5. Use Activate / Deactivate to block a customer from logging in without deleting their order history.",
    "[Screenshot: A3 — Manage customers]",
    "",
    "3E. Manage orders — update status",
    "1. Click Orders in the sidebar.",
    "2. Select an order from the list.",
    "3. View line items and total in the detail area.",
    "4. Choose a new status: Pending, Ready for Pickup, or Delivered.",
    "5. Click Update Status.",
    "6. For orders with prescriptions, open the prescription file and click Verify or Reject before releasing medicine.",
    "[Screenshot: A4 — Manage orders]",
    "",
    "3F. Reports",
    "1. Click Reports in the sidebar.",
    "2. Switch between Sales, Stock, and Customer history tabs (if available).",
    "3. Read summary figures for management decisions.",
    "4. Export or print if your build supports it.",
    "[Screenshot: A5 — Reports]",
    "",
    "PART 4 — Customer user guide",
    "",
    "4A. Customer home",
    "1. After customer login, Home shows a welcome summary.",
    "2. Note promotions or cart reminders.",
    "3. Use Browse to shop or Orders to track purchases.",
    "[Screenshot: C1 — Customer home]",
    "",
    "4B. Search and browse medicines",
    "1. Click Browse in the sidebar.",
    "2. Type part of a medicine name in the search box, or pick a category and price range.",
    "3. Click Search or press Enter.",
    "4. Click a row or Details to read full product information.",
    "5. Enter quantity and click Add to Cart.",
    "[Screenshot: C2 — Browse medicines]",
    "[Screenshot: C3 — Medicine details]",
    "",
    "4C. Cart and checkout",
    "1. Click Cart in the sidebar.",
    "2. Review each line: name, quantity, price, and total.",
    "3. For prescription (Rx) items, click the prescription column and upload a PDF or image file for each Rx line.",
    "4. Tick which items you want to buy if some lines are unchecked.",
    "5. Click Place Order.",
    "6. In the payment dialog, choose payment method and confirm the total.",
    "7. Click Confirm. A success message shows your order number.",
    "[Screenshot: C4 — Cart]",
    "[Screenshot: C5 — Payment dialog]",
    "",
    "4D. Track orders",
    "1. Click Orders in the sidebar.",
    "2. Select your order from the list.",
    "3. Read status: Pending, Ready for Pickup, Delivered, or Cancelled.",
    "4. View line items and amounts.",
    "5. To cancel, select a Pending order, click Cancel, and enter a reason.",
    "6. To export history, use Export PDF or Export CSV if shown.",
    "[Screenshot: C6 — Track orders]",
    "",
    "4E. Update profile",
    "1. Click Profile in the sidebar.",
    "2. Edit name, phone, or address.",
    "3. Click Save.",
    "4. Changes apply to future orders and contact details.",
    "[Screenshot: C7 — Profile]",
    "",
    "PART 5 — Troubleshooting",
    "• Login fails — check username/email and password; customers must register first.",
    "• Checkout fails — check stock, upload all required prescriptions, and ensure medicines are still available.",
    "• Admin cannot deactivate medicine — order still pending or ready; complete or cancel it first.",
    "• Database error — ask administrator to check SQL Server and App.config connection string.",
]

CLASS_DETAILS: dict[str, str] = {
    "User": (
        "The User class is the parent type for people who can sign in to SmartMed. It represents shared account information "
        "that both administrators and customers need, such as name, contact details, and login credentials.\n\n"
        "In object-oriented terms, this class demonstrates generalisation: common data is defined once and reused by "
        "specialised child classes. The models are kept as simple data containers (properties only). Methods such as login "
        "or profile update are not placed inside User because those actions involve database access, validation across "
        "several fields, and navigation between screens. That behaviour belongs in AuthService and CustomerService, which "
        "is a standard layered approach for maintainability and matches the three-tier architecture."
    ),
    "Admin": (
        "Admin extends User and represents a pharmacy staff member who manages the system. Inheritance allows Admin to "
        "reuse all User fields while adding an administrator identifier and username.\n\n"
        "No business methods are declared on Admin itself. Opening the dashboard or managing stock are full application "
        "workflows, not single-object operations, so they are implemented in forms and services. After login, an Admin "
        "object is stored in the session service so the rest of the admin screens know who is signed in. This is encapsulation "
        "at application level: the UI works with a typed object without handling raw database rows."
    ),
    "Customer": (
        "Customer extends User and models a registered pharmacy client. Inheritance avoids duplicating name, email, phone, "
        "and password on a separate unrelated class.\n\n"
        "As with Admin, Customer does not contain methods like PlaceOrder or TrackOrder. Those use cases coordinate the cart, "
        "stock, prescriptions, and database transactions. Placing that logic in OrderService and related classes follows "
        "the Single Responsibility idea: the Customer class describes who the user is; services describe what the system does "
        "for them. The IsActive flag on the user record supports account deactivation without deleting history."
    ),
    "Medicine": (
        "Medicine is a core entity class for products in the pharmacy catalogue. It groups everything the application needs "
        "to know about a item for sale — identity, pricing, stock, expiry, promotions, and whether a prescription is required.\n\n"
        "The class uses encapsulation through public properties that map to database columns. Behaviour such as checking expiry, "
        "applying a promotion price, or deciding if an item can be sold is handled in MedicineService because those rules may "
        "combine several fields and must be applied consistently on admin screens, search, cart, and checkout. Keeping Medicine "
        "as a data class makes repositories straightforward: they load and save records without embedding pharmacy policy."
    ),
    "Order": (
        "Order represents one customer purchase as a whole. It is an association class in the domain model linking a customer "
        "to a point in time, a status, and a total amount.\n\n"
        "Order does not include methods like CalculateTotal in the model layer because totals are produced when the order is "
        "built in OrderService from line items and current prices. Status changes also pass through service and repository "
        "layers so rules (for example, when cancellation is allowed) are enforced in one place. This separation keeps the entity "
        "easy to bind to grids and reports while preserving business control in services."
    ),
    "OrderItem": (
        "OrderItem links an order to one medicine line with quantity and price information captured at purchase time. "
        "It supports the one-to-many relationship between orders and products in the database design.\n\n"
        "Storing unit price and subtotal on each line is an important design choice: if catalogue prices change later, "
        "historical orders still show what the customer actually paid. The class remains property-based; calculation and "
        "validation happen when OrderService creates items during checkout."
    ),
    "Prescription": (
        "Prescription models an uploaded prescription file tied to a customer order, and optionally to a specific medicine "
        "when more than one prescription-required product is in the same checkout.\n\n"
        "File copying, status workflow (pending, verified, rejected), and admin approval are implemented in PrescriptionService "
        "rather than in the entity class, because they involve the file system and cross-record updates. The Prescription class "
        "therefore supports encapsulation of data while the service layer encapsulates the verification process."
    ),
}

SERVICE_SECTION = (
    "Service classes group the main business behaviour of SmartMed. They sit between forms and repositories.\n\n"
    "AuthService handles login, registration, and password changes. MedicineService manages the catalogue, stock rules, "
    "promotions, and expiry checks. OrderService and CartService handle shopping and checkout. PrescriptionService manages "
    "uploaded prescription files and admin approval. ReportService prepares dashboard and report figures. ValidationService "
    "and PasswordHasher provide shared technical support.\n\n"
    "This layout applies object-oriented design at module level: each service has a focused responsibility, collaborates with "
    "repositories (composition), and works with model objects passed as parameters. Forms depend on services, not on SQL, "
    "which keeps the presentation layer thin and easier to change."
)

DIAGRAM_IMAGES = [
    ("Architecture Diagram", "Architecture.png"),
    ("Use Case Diagram", "UseCase.png"),
    ("ER Diagram", "ER.png"),
    ("Class Diagram", "Class.png"),
    ("Sequence Diagram", "Sequence.png"),
]

SKIP_CLEAR_HEADINGS = frozenset({"Functional Requirements"})


def set_body_format(paragraph):
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def find_heading_index(doc, title, level_prefix="Heading"):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == title and p.style.name.startswith(level_prefix):
            return i
    return None


def remove_paragraph(p):
    p._element.getparent().remove(p._element)


def clear_body_between(doc, heading_title, heading_prefix="Heading"):
    idx = find_heading_index(doc, heading_title, heading_prefix)
    if idx is None:
        return False
    level = int(doc.paragraphs[idx].style.name.split()[-1]) if doc.paragraphs[idx].style.name.split()[-1].isdigit() else 1
    to_remove = []
    for j in range(idx + 1, len(doc.paragraphs)):
        para = doc.paragraphs[j]
        if para.style.name.startswith("Heading"):
            try:
                next_level = int(para.style.name.split()[-1])
            except ValueError:
                next_level = 1
            if next_level <= level:
                break
        to_remove.append(para)
    for p in to_remove:
        remove_paragraph(p)
    return True


def insert_bodies_after_heading(doc, heading_title, bodies, style="Body Text"):
    idx = find_heading_index(doc, heading_title)
    if idx is None:
        print(f"Heading not found: {heading_title}")
        return
    anchor = doc.paragraphs[idx]
    for body in bodies:
        if body == "":
            continue
        new_p = insert_paragraph_after(anchor)
        try:
            new_p.style = doc.styles[style]
        except KeyError:
            pass
        new_p.add_run(body)
        set_body_format(new_p)
        anchor = new_p
    print(f"Filled: {heading_title} ({len(bodies)} paragraphs)")


def insert_list_items_after_heading(doc, heading_title, items, use_list_style=False):
    idx = find_heading_index(doc, heading_title, "Heading 3")
    if idx is None:
        return
    clear_body_between(doc, heading_title, "Heading 3")
    idx = find_heading_index(doc, heading_title, "Heading 3")
    anchor = doc.paragraphs[idx]
    for item in items:
        new_p = insert_paragraph_after(anchor)
        if use_list_style:
            try:
                new_p.style = doc.styles["List Paragraph"]
            except KeyError:
                pass
        else:
            try:
                new_p.style = doc.styles["Body Text"]
            except KeyError:
                pass
        new_p.add_run(item)
        set_body_format(new_p)
        anchor = new_p
    print(f"Restored functional list: {heading_title}")


def restore_functional_requirements(doc):
    fr_idx = find_heading_index(doc, "Functional Requirements", "Heading 2")
    if fr_idx is None:
        return

    def fill_fr_subsection(title, items, list_style):
        idx = None
        for i in range(fr_idx + 1, len(doc.paragraphs)):
            p = doc.paragraphs[i]
            if p.style.name == "Heading 2" and p.text.strip() not in ("", "Functional Requirements"):
                if idx is not None:
                    break
            if p.text.strip() == title and p.style.name == "Heading 3":
                idx = i
                break
        if idx is None:
            return
        to_remove = []
        for j in range(idx + 1, len(doc.paragraphs)):
            para = doc.paragraphs[j]
            if para.style.name == "Heading 3" or (para.style.name == "Heading 2" and para.text.strip() != ""):
                break
            to_remove.append(para)
        for p in to_remove:
            remove_paragraph(p)
        idx = None
        for i in range(fr_idx + 1, len(doc.paragraphs)):
            p = doc.paragraphs[i]
            if p.text.strip() == title and p.style.name == "Heading 3":
                idx = i
                break
        anchor = doc.paragraphs[idx]
        for item in items:
            new_p = insert_paragraph_after(anchor)
            if list_style:
                try:
                    new_p.style = doc.styles["List Paragraph"]
                except KeyError:
                    pass
            else:
                try:
                    new_p.style = doc.styles["Body Text"]
                except KeyError:
                    pass
            new_p.add_run(item)
            set_body_format(new_p)
            anchor = new_p
        print(f"Restored functional list: {title}")

    fill_fr_subsection("Admin", FUNCTIONAL_ADMIN, list_style=False)
    fill_fr_subsection("Customer", FUNCTIONAL_CUSTOMER, list_style=False)
    fill_fr_subsection("Additional Features", FUNCTIONAL_ADDITIONAL, list_style=True)


def paragraph_has_image(paragraph):
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            return True
    return False


def insert_image_after_heading(doc, heading_title, image_path, caption=None, width=6.0):
    if not os.path.isfile(image_path):
        print(f"Image missing: {image_path}")
        return False
    idx = find_heading_index(doc, heading_title)
    if idx is None:
        return False
    for j in range(idx + 1, min(idx + 8, len(doc.paragraphs))):
        if doc.paragraphs[j].style.name.startswith("Heading"):
            break
        if paragraph_has_image(doc.paragraphs[j]):
            print(f"Image already present: {heading_title}")
            return True
    anchor = doc.paragraphs[idx]
    img_p = insert_paragraph_after(anchor)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(image_path, width=Inches(width))
    if caption:
        cap = insert_paragraph_after(img_p)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.italic = True
    print(f"Inserted diagram: {heading_title}")
    return True


def update_class_descriptions(doc):
    classes_start = find_heading_index(doc, "Description of Classes, Properties and Methods")
    if classes_start is None:
        print("Classes section not found")
        return

    for class_name, body in CLASS_DETAILS.items():
        idx = None
        for i in range(classes_start + 1, len(doc.paragraphs)):
            p = doc.paragraphs[i]
            if p.text.strip() == class_name and p.style.name == "Heading 3":
                idx = i
                break
        if idx is None:
            continue

        level = 3
        to_remove = []
        for j in range(idx + 1, len(doc.paragraphs)):
            para = doc.paragraphs[j]
            if para.style.name.startswith("Heading"):
                try:
                    next_level = int(para.style.name.split()[-1])
                except ValueError:
                    next_level = 1
                if next_level <= level:
                    break
            to_remove.append(para)
        for p in to_remove:
            remove_paragraph(p)

        idx = None
        for i in range(classes_start + 1, len(doc.paragraphs)):
            p = doc.paragraphs[i]
            if p.text.strip() == class_name and p.style.name == "Heading 3":
                idx = i
                break
        if idx is None:
            continue

        anchor = doc.paragraphs[idx]
        for part in body.split("\n\n"):
            new_p = insert_paragraph_after(anchor)
            try:
                new_p.style = doc.styles["Body Text"]
            except KeyError:
                pass
            new_p.add_run(part)
            set_body_format(new_p)
            anchor = new_p
        print(f"Updated class: {class_name}")


def insert_heading3_after(anchor, doc, title):
    h = insert_paragraph_after(anchor)
    h.style = doc.styles["Heading 3"]
    h.add_run(title)
    return h


def insert_bold_label_after(anchor, doc, text):
    new_p = insert_paragraph_after(anchor)
    try:
        new_p.style = doc.styles["Body Text"]
    except KeyError:
        pass
    r = new_p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    new_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return new_p


def append_body_paragraphs(anchor, doc, texts):
    for text in texts:
        new_p = insert_paragraph_after(anchor)
        try:
            new_p.style = doc.styles["Body Text"]
        except KeyError:
            pass
        new_p.add_run(text)
        set_body_format(new_p)
        anchor = new_p
    return anchor


def set_heading_text(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def remove_heading_section(doc, heading_title, heading_prefix="Heading"):
    """Remove heading paragraph and all body until the next heading of equal or higher level."""
    idx = find_heading_index(doc, heading_title, heading_prefix)
    if idx is None:
        return False
    level = int(doc.paragraphs[idx].style.name.split()[-1]) if doc.paragraphs[idx].style.name.split()[-1].isdigit() else 1
    to_remove = [doc.paragraphs[idx]]
    for j in range(idx + 1, len(doc.paragraphs)):
        para = doc.paragraphs[j]
        if para.style.name.startswith("Heading"):
            try:
                next_level = int(para.style.name.split()[-1])
            except ValueError:
                next_level = 1
            if next_level <= level:
                break
        to_remove.append(para)
    for p in to_remove:
        remove_paragraph(p)
    print(f"Removed section: {heading_title}")
    return True


def merge_implementation_headings(doc):
    """Keep one Implementation H2; remove separate Code Snippets and Output Screens sections."""
    merged_idx = find_heading_index(doc, IMPLEMENTATION_HEADING)
    old_merged = find_heading_index(
        doc, "Brief Description of Functions, Code Snippets and Output Screens"
    )
    anchor_idx = merged_idx if merged_idx is not None else old_merged

    if anchor_idx is not None:
        set_heading_text(doc.paragraphs[anchor_idx], IMPLEMENTATION_HEADING)

    for title in ("Output Screens", "Code Snippets",
                  "Brief Description of Functions, Code Snippets and Output Screens"):
        while remove_heading_section(doc, title):
            pass

    print(f"Implementation section: {IMPLEMENTATION_HEADING}")


def find_implementation_heading(doc):
    return find_heading_index(doc, IMPLEMENTATION_HEADING)


def fill_implementation_section(doc):
    merge_implementation_headings(doc)
    idx = find_implementation_heading(doc)
    if idx is None:
        print("Implementation heading not found")
        return
    title = doc.paragraphs[idx].text.strip()
    clear_body_between(doc, title)
    idx = find_implementation_heading(doc)
    anchor = doc.paragraphs[idx]

    for block_title, paragraphs, code_snippets, screenshot in FEATURE_BLOCKS:
        anchor = insert_heading3_after(anchor, doc, block_title)
        anchor = append_body_paragraphs(anchor, doc, paragraphs)
        if code_snippets:
            anchor = insert_bold_label_after(anchor, doc, "Code snippets to attach:")
            anchor = append_body_paragraphs(
                anchor, doc, [f"• {item}" for item in code_snippets]
            )
        if screenshot:
            anchor = insert_bold_label_after(anchor, doc, "Interface screenshot:")
            anchor = append_body_paragraphs(
                anchor, doc, [screenshot, "[Insert screenshot of this interface here]"]
            )

    print("Filled: Implementation (description + snippets + screenshots)")


def fill_user_manual(doc):
    idx = find_heading_index(doc, "Using the Application")
    if idx is None:
        return
    clear_body_between(doc, "Using the Application")
    idx = find_heading_index(doc, "Using the Application")
    anchor = doc.paragraphs[idx]
    for text in USER_MANUAL_PARAS:
        if text == "":
            continue
        new_p = insert_paragraph_after(anchor)
        try:
            new_p.style = doc.styles["Body Text"]
        except KeyError:
            pass
        new_p.add_run(text)
        set_body_format(new_p)
        anchor = new_p
    print("Filled: Using the Application (expanded user manual)")


def insert_service_section(doc):
    idx = find_heading_index(doc, "Service Classes", "Heading 3")
    if idx is not None:
        clear_body_between(doc, "Service Classes", "Heading 3")
        idx = find_heading_index(doc, "Service Classes", "Heading 3")
        anchor = doc.paragraphs[idx]
        p = insert_paragraph_after(anchor)
        try:
            p.style = doc.styles["Body Text"]
        except KeyError:
            pass
        p.add_run(SERVICE_SECTION)
        set_body_format(p)
        print("Updated Service Classes")
        return
    idx = find_heading_index(doc, "Prescription", "Heading 3")
    impl_idx = find_heading_index(doc, "Implementation")
    if idx is None or impl_idx is None:
        return
    anchor = doc.paragraphs[impl_idx - 1]
    h = insert_paragraph_after(anchor)
    h.style = doc.styles["Heading 3"]
    h.add_run("Service Classes")
    p = insert_paragraph_after(h)
    try:
        p.style = doc.styles["Body Text"]
    except KeyError:
        pass
    p.add_run(SERVICE_SECTION)
    set_body_format(p)
    print("Inserted Service Classes section")


def fill_report():
    if not os.path.isfile(REPORT_PATH):
        raise FileNotFoundError(REPORT_PATH)

    doc = Document(REPORT_PATH)

    restore_functional_requirements(doc)

    skip = {
        IMPLEMENTATION_HEADING,
        "Brief Description of Functions",
        "Using the Application",
    }
    for heading, bodies in FILL_AFTER_HEADING.items():
        if heading in skip:
            continue
        clear_body_between(doc, heading)
        insert_bodies_after_heading(doc, heading, bodies)

    fill_implementation_section(doc)
    fill_user_manual(doc)

    update_class_descriptions(doc)
    insert_service_section(doc)

    restore_functional_requirements(doc)  # ensure FR not touched by class pass

    for heading, filename in DIAGRAM_IMAGES:
        path = os.path.join(EXPORT_DIR, filename)
        insert_image_after_heading(doc, heading, path, caption=f"Figure: {heading}")

    er_path = os.path.join(EXPORT_DIR, "ER.png")
    insert_image_after_heading(doc, "Database Design", er_path, caption="Figure: Database ER model", width=5.5)

    tmp = REPORT_PATH + ".tmp"
    doc.save(tmp)
    try:
        os.replace(tmp, REPORT_PATH)
        print(f"Saved {REPORT_PATH}")
    except OSError:
        fallback = os.path.join(os.path.dirname(REPORT_PATH), "SmartMed Report Sample (merged).docx")
        try:
            doc.save(fallback)
            if os.path.exists(tmp):
                os.remove(tmp)
            print(f"Original file is open in Word. Saved merged copy to:\n  {fallback}", file=sys.stderr)
        except OSError as ex:
            if os.path.exists(tmp):
                os.remove(tmp)
            print(f"Close Word and run again. ({ex})", file=sys.stderr)
            sys.exit(1)


if __name__ == "__main__":
    fill_report()
