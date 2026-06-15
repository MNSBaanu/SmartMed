"""Fill SmartMed Report.docx with implementation documentation."""
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING
import os

REPORT_PATH = os.path.normpath(
    os.path.join(os.path.dirname(__file__), "..", "Report", "SmartMed Report.docx")
)

SECTIONS = {
    "Introduction": """SmartMed Pharmacy is a desktop Windows Forms application developed in C# using the .NET Framework 4.8. The system supports two user roles: administrators who manage inventory, customers, and orders, and customers who search for medicines, place orders, and track fulfilment status. The application follows a three-tier architecture separating presentation (SmartMed.UI), business logic (SmartMed.Business), and data access (SmartMed.Data) with persistence in Microsoft SQL Server.

This report documents how to run the software, describes the architecture and class design aligned with the coursework diagrams, explains the search algorithms used for medicine catalogue queries, and includes a reflective essay on the development experience. The implementation satisfies the mandatory functional requirements from the assignment brief, including login, registration, medicine CRUD, customer management, order processing, reports, dashboard, and linear search with filtering.""",

    "Detailed Instructions to Run the Program": """Prerequisites:
• Windows 10 or later with .NET Framework 4.8 installed.
• Visual Studio 2019 or later with the .NET desktop development workload.
• SQL Server Express, LocalDB, or full SQL Server with Windows Authentication.

Database setup:
1. Open SQL Server Management Studio (SSMS).
2. Open the script Database/SmartMedDB.sql from the project folder.
3. Execute the script. It creates the SmartMedDB database, all tables, constraints, and seed data.
4. Default admin credentials: Username admin, Password admin123.
5. Sample customer: john@email.com / customer123.

Connection string:
1. Open SmartMed.UI/App.config.
2. Adjust the SmartMedDB connection string if your SQL Server instance is not the default local instance (.). For LocalDB use: Data Source=(localdb)\\MSSQLLocalDB;Initial Catalog=SmartMedDB;Integrated Security=True

Build and run:
1. Open SmartMed.sln in Visual Studio.
2. Right-click SmartMed.UI and Set as Startup Project.
3. Build → Build Solution (Ctrl+Shift+B). Alternatively run: dotnet build SmartMed.sln from the solution folder.
4. Press F5 to start debugging. The Login form appears first.
5. Log in as admin (admin / admin123) to access the admin dashboard, or register a new customer account.

Admin workflow: Dashboard overview → Manage Medicines → Manage Customers → Manage Orders → Reports.
Customer workflow: Register or login → Search Medicines → Place Order (cart) → Track Orders → Profile.""",

    "System Overview": """The SmartMed system is a single-user-session desktop application. On startup, Program.cs launches LoginForm. Successful authentication stores the current user in Session (SmartMed.Business) and opens either AdminDashboardForm or CustomerDashboardForm based on role. Each dashboard provides navigation buttons to feature-specific forms. All database access is routed through repository classes in SmartMed.Data; WinForms code never opens SQL connections directly, preserving separation of concerns.""",

    "Three-Layer Architecture": """Presentation Layer (SmartMed.UI): Windows Forms for login, registration, admin dashboards, medicine CRUD, customer management, order management, reports, medicine search, cart/checkout, order tracking, and profile editing.

Business Layer (SmartMed.Business): Entity classes (Person, Customer, Admin, Medicine, Order, OrderItem, Prescription), ValidationHelper for input rules, SearchHelper for linear search and filtering, Session for authenticated user state, and service classes (AuthService, MedicineService, OrderService, CustomerService, DashboardService) that coordinate validation and repository calls.

Data Access Layer (SmartMed.Data): DatabaseHelper wraps SqlConnection and parameterized commands. Repositories (AdminRepository, CustomerRepository, MedicineRepository, OrderRepository) implement CRUD and queries. Data models map to database rows.

Database Layer: SQL Server database SmartMedDB with tables Admin, Customer, Medicine, Order, OrderItem, and Prescription. Foreign keys link orders to customers and order items to orders and medicines. Order.Status is constrained to Pending, Ready for Pickup, or Delivered.""",

    "Use Case Diagram": """The use case diagram (Diagrams/Usecase.drawio) identifies Admin actors (login, manage medicines, manage customers, manage orders, reports, dashboard) and Customer actors (register, login, search medicines, place orders, track orders, manage profile). Include relationships show that placing an order includes searching medicines and that order management includes status updates. The implemented forms map directly to these use cases.""",

    "Entity-Relationship Diagram": """The ER diagram (Diagrams/ER.drawio) models six entities. Admin is independent. Customer places Orders (one-to-many). Order contains OrderItems (one-to-many). Medicine appears in OrderItems (many-to-many resolved via OrderItem). Customer uploads Prescriptions (one-to-many). Primary keys are identity integers. The SQL script Database/SmartMedDB.sql implements this schema with appropriate foreign keys and a CHECK constraint on order status.""",

    "Class Diagram": """The class diagram (Diagrams/Class.drawio) shows Person as an abstract base with Name, Email, Phone, and Password. Customer and Admin inherit from Person. Customer adds CustomerID, Address, Register(), and Login(). Admin adds AdminID, Username, and Login(). Medicine, Order, OrderItem, and Prescription are separate entity classes. Business layer classes mirror this design; data layer uses flat DTO-style models for efficient ADO.NET mapping.""",

    "Class Relationships": """Inheritance: Customer and Admin extend Person, enabling shared validation of email and password fields. Composition: Order aggregates OrderItem records; each OrderItem references a Medicine. Association: Order belongs to Customer; Prescription belongs to Customer. Services depend on repositories (dependency direction: Business → Data). UI forms depend on services only, not repositories, maintaining the vertical architecture shown in Diagrams/Architecture.drawio.""",

    "Source Code Attribution": """All source code in this submission was written for the SmartMed coursework. Standard Microsoft .NET Framework libraries (System.Windows.Forms, System.Data.SqlClient, System.Configuration) are used as documented. No third-party UI frameworks were used. Diagram source files are original draw.io diagrams created for this module. Any external references (Microsoft documentation for ADO.NET and WinForms) informed API usage but were not copied as code.""",

    "Admin Features": """Login: AuthService validates admin username and password against the Admin table via AdminRepository. Empty fields are rejected by ValidationHelper before database access.

Manage Medicines: MedicineManagementForm provides add, update, delete, and list operations. Fields include name, category, dosage, price, stock quantity, supplier, expiry date, and prescription requirement. Numeric validation ensures price and stock are valid.

Manage Customers: CustomerManagementForm lists registered customers and allows administrators to update contact details.

Manage Orders: OrderManagementForm displays all orders with customer name, date, status, and total. Administrators can change status to Pending, Ready for Pickup, or Delivered.

Reports: ReportsForm shows sales totals, stock summary, and customer order history using queries from OrderRepository and MedicineRepository.

Dashboard: AdminOverviewForm displays aggregate metrics (total sales, medicines in stock, active orders) via DashboardService.""",

    "Customer Features": """Registration: RegistrationForm collects full name, email, phone, address, password, and confirmation. ValidationHelper checks required fields, email format, and password match before CustomerRepository inserts a new row.

Login: Customers authenticate with email and password; Session stores the logged-in Customer object.

Search Medicines: SearchMedicinesForm loads the catalogue and applies SearchHelper linear search by name, filter by category, and filter by price range. Results bind to a DataGridView.

Place Orders: PlaceOrderForm lists available medicines, supports quantity entry, maintains an in-memory cart, and calls OrderService.PlaceOrder to insert Order and OrderItem rows and decrement stock.

Track Orders: TrackOrdersForm lists the current customer's orders and statuses.

Profile Management: ProfileManagementForm allows customers to update their contact information.""",

    "Additional Features": """Optional features from the brief (discounts, expiry notifications, prescription upload UI, PDF/Excel export) were deferred to prioritise mandatory marking-scheme tasks. The Prescription table and Prescription business class exist in the schema and model layer for future extension. Core order and inventory workflows are fully functional without these extras.""",

    "Programming Language and Platform": """The application is implemented in C# targeting .NET Framework 4.8. The UI project uses Windows Forms (System.Windows.Forms). The solution uses SDK-style projects compatible with Visual Studio 2019+ and the dotnet CLI build. This meets the requirement for Visual Studio 2015 or higher compatibility at the language and framework level.""",

    "Data Storage": """Persistent data is stored in Microsoft SQL Server. The connection string named SmartMedDB in App.config is read by DatabaseHelper using ConfigurationManager. Parameterized SQL prevents injection. Transactions are used when placing orders to ensure order header, line items, and stock updates succeed or roll back together.""",

    "Software Design": """Object-oriented design uses inheritance (Person hierarchy), encapsulation (private repository calls inside services), and separation of layers. Each form handles only UI events; business rules live in services and helpers. Repositories isolate SQL syntax from the rest of the application, simplifying maintenance and testing.""",

    "User Interface": """Forms are built programmatically in code for portability (no designer files required). A consistent colour scheme uses blue tones matching project diagrams. Role-based navigation prevents customers from accessing admin forms. MessageBox displays validation errors and database exceptions in user-friendly language.""",

    "Validation and Exception Handling": """ValidationHelper provides IsNullOrEmpty, IsValidEmail, IsPositiveDecimal, and IsPositiveInteger checks used across login, registration, and medicine forms. Service and repository methods wrap database operations in try/catch blocks; SQLException messages are surfaced via MessageBox on the UI thread. Order placement validates stock availability before committing.""",

    "Person Class": """Person (SmartMed.Business.Models.Person) is the base class with properties: Name, Email, Phone, Password. It represents shared attributes for human users. Derived classes add role-specific identifiers and behaviour.""",

    "Customer Class": """Customer extends Person with CustomerID (int) and Address (string). Methods Register() and Login() are documented in the class diagram; the implemented workflow delegates to AuthService and CustomerRepository while Session holds the authenticated instance after login.""",

    "Admin Class": """Admin extends Person with AdminID (int) and Username (string). Admin login uses username rather than email. AuthService.LoginAdmin queries AdminRepository and populates Session.CurrentAdmin on success.""",

    "Medicine Class": """Medicine encapsulates MedicineID, MedicineName, Category, Dosage, Price, StockQuantity, Supplier, ExpiryDate, and RequiresPrescription. MedicineService wraps MedicineRepository for CRUD used by admin and search features.""",

    "Order Class": """Order represents OrderID, CustomerID, OrderDate, Status, and TotalAmount. OrderService creates orders, lists by customer, and updates status for admin fulfilment.""",

    "OrderItem Class": """OrderItem links an order to a medicine with Quantity, UnitPrice, and Subtotal. Subtotal is calculated as Quantity × UnitPrice at the time of order placement to preserve historical pricing.""",

    "Prescription Class": """Prescription models PrescriptionID, CustomerID, PrescriptionFile path, UploadDate, and Status. The database table is created; upload UI is reserved for future enhancement.""",

    "Search Algorithms Used in the Project": """Medicine search uses linear search and filtering as required by the coursework. SearchHelper.SearchByName iterates the full medicine list once (O(n)), comparing each MedicineName to the keyword using case-insensitive Contains. FilterByCategory applies a second linear pass when a category filter is active. FilterByPriceRange iterates and retains items where Price falls between min and max inclusive.

The composite Search method chains these operations: start with the full list, optionally narrow by name, then category, then price range. Each step produces a new filtered list. This is documented as linear search with filtering rather than binary search, which requires sorted data and was not necessary for the catalogue size.

Pseudocode:
  results = all medicines
  if name provided: results = LinearSearchByName(results, name)
  if category provided: results = LinearFilterByCategory(results, category)
  if price range provided: results = LinearFilterByPrice(results, min, max)
  return results

SearchMedicinesForm calls MedicineService.Search, which loads medicines from the database and delegates to SearchHelper.Search before binding results to the grid.""",

    "Experience with C# and Visual Studio": """Developing SmartMed strengthened my understanding of C# as an object-oriented language and of Visual Studio as an integrated environment for building desktop applications. Before this module my experience was limited to console programs; WinForms introduced event-driven programming where button clicks and form load events trigger methods that update the interface asynchronously from user input.

Working with a multi-project solution clarified how references work: SmartMed.UI references Business, which references Data. Building the solution compiles dependencies in order. Debugging with breakpoints across layers showed how a click on Login propagates from LoginForm through AuthService to CustomerRepository and back. The Immediate Window and Watch panel helped inspect Session state during authentication.

I used both Visual Studio and the dotnet CLI to build. SDK-style projects simplified the csproj files compared to older templates. Configuring App.config for the connection string reinforced the importance of externalising environment-specific settings rather than hard-coding server names in source code.""",

    "Features Liked and Rationale": """I particularly valued the three-tier architecture because it mirrors professional practice. When I needed to change how medicines were loaded, I edited MedicineRepository once while MedicineManagementForm and SearchMedicinesForm continued to call MedicineService unchanged. This separation reduced duplication and made the codebase easier to reason about.

Windows Forms was approachable for rapid layout of grids, text boxes, and buttons. DataGridView data binding to lists of MedicineItem provided instant tabular display without manual row painting. The inheritance model for Person, Customer, and Admin made the domain model intuitive and aligned with the class diagram produced in the design phase.

Implementing SearchHelper as a static utility class kept search logic testable and documented independently of the UI. I liked that linear search is easy to explain in documentation and sufficient for pharmacy catalogues that are not millions of rows.""",

    "Challenges Faced and Solutions Applied": """A significant challenge was ensuring the data layer could read the connection string when called from a class library. DatabaseHelper uses ConfigurationManager, which reads from the executing assembly's config file. The UI project's App.config is copied to SmartMed.exe.config on build, so running the executable provides the correct connection string to all layers.

Order placement required transactional integrity: inserting an order, multiple order items, and updating stock had to succeed together. OrderRepository wraps these steps in a SqlTransaction so a failure mid-process rolls back partial changes.

Another challenge was keeping WinForms code-behind thin while still showing helpful errors. I standardised on try/catch in event handlers that display ex.Message in MessageBox, while repositories throw descriptive exceptions for constraint violations.

Building forms entirely in code (without the designer) was initially verbose but avoided merge conflicts and kept the repository self-contained for submission. I reused patterns for grid setup, button placement, and BackColor theming across admin and customer forms.

Mapping between Business models and Data models (e.g. Customer vs CustomerUser) required discipline. Services translate between layers when needed, keeping ADO.NET details out of the business entities.""",

    "Future Improvements": """Future versions could add prescription file upload with validation, email notifications when order status changes, and export of reports to PDF or Excel using libraries such as iTextSharp or ClosedXML. Role-based admin permissions could distinguish pharmacists from managers. A barcode scanner integration would speed stock intake. For larger catalogues, indexing and full-text search in SQL Server could supplement client-side linear search. Unit tests for ValidationHelper and SearchHelper would improve regression safety. Migrating to .NET 6+ WinForms would enable cross-platform considerations, though deployment targets remain Windows for pharmacy workstations.""",

    "Conclusion": """SmartMed delivers a working pharmacy management desktop application that meets the coursework functional and technical requirements. The design diagrams guided implementation of the database schema, class structure, and layered architecture. Mandatory features—authentication, medicine CRUD, customer search with linear filtering, order lifecycle, reports, and dashboard—are implemented and runnable against SQL Server seed data. The project improved my skills in C#, WinForms, ADO.NET, and structured OOP design. Documentation and reflective analysis consolidate learning outcomes for assessment.""",

    "References": """Microsoft (2024). ADO.NET documentation. https://learn.microsoft.com/en-us/dotnet/framework/data/adonet/

Microsoft (2024). Windows Forms documentation. https://learn.microsoft.com/en-us/dotnet/desktop/winforms/

Module coursework brief: Document.md (SmartMed Pharmacy specification).""",

    "Appendix A: Use Case Diagram": """See Diagrams/Usecase.drawio — export as PNG for printed submission if required.""",

    "Appendix B: Entity-Relationship Diagram": """See Diagrams/ER.drawio — Chen notation with Admin, Customer, Medicine, Order, OrderItem, Prescription.""",

    "Appendix C: Class Diagram": """See Diagrams/Class.drawio — Person inheritance, entity classes, and relationships.""",

    "Appendix D: Architecture Diagram": """See Diagrams/Architecture.drawio — vertical flow: Presentation → Business → Data Access → SQL Server.""",
}

HEADINGS = set(SECTIONS.keys())


def insert_paragraph_after(paragraph, text=""):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def set_paragraph_format(paragraph):
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def is_section_heading(text):
    return text.strip() in HEADINGS


def fill_report():
    doc = Document(REPORT_PATH)
    idx = 0
    while idx < len(doc.paragraphs):
        title = doc.paragraphs[idx].text.strip()
        if title not in SECTIONS:
            idx += 1
            continue

        # Remove existing body paragraphs until the next heading
        while idx + 1 < len(doc.paragraphs):
            nxt = doc.paragraphs[idx + 1].text.strip()
            if nxt in HEADINGS:
                break
            rm = doc.paragraphs[idx + 1]
            rm._element.getparent().remove(rm._element)

        anchor = doc.paragraphs[idx]
        last = anchor
        for part in SECTIONS[title].split("\n\n"):
            new_p = insert_paragraph_after(last, part)
            set_paragraph_format(new_p)
            last = new_p
        idx += 1

    doc.save(REPORT_PATH)
    word_count = sum(len(v.split()) for v in SECTIONS.values())
    print(f"Saved {REPORT_PATH}")
    print(f"Approximate word count: {word_count}")


if __name__ == "__main__":
    fill_report()
